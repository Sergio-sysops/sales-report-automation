"""Turn a messy sales CSV into a clean, formatted Excel report with charts.

Usage:
    python sales_report.py ventas_sucias.csv -o sales_report.xlsx
    python sales_report.py ventas_sucias.csv --email boss@example.com   # needs SMTP_* env vars
"""
import argparse
import os
import smtplib
from email.message import EmailMessage

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREEN = "1A9E5C"
DARK = "0A0F16"
LIGHT = "EAF7F0"
MONEY = '#,##0.00 "€"'


DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%b %d %Y"]


def parse_dates(col):
    """Try each known format explicitly, so 2026-03-10 never becomes 3 October."""
    result = pd.Series(pd.NaT, index=col.index)
    for fmt in DATE_FORMATS:
        result = result.fillna(pd.to_datetime(col, format=fmt, errors="coerce"))
    return result


def clean(df):
    stats = {"rows_in": len(df)}
    df = df.dropna(how="all")
    stats["empty_removed"] = stats["rows_in"] - len(df)

    before = len(df)
    df = df.drop_duplicates()
    stats["duplicates_removed"] = before - len(df)

    df["date"] = parse_dates(df["date"])
    df["product"] = df["product"].str.split().str.join(" ").str.title().str.replace("Usb-C", "USB-C").str.replace("Hd", "HD")
    df["region"] = df["region"].fillna("Unknown").replace("", "Unknown").str.title()
    df["seller"] = df["seller"].fillna("Unknown").replace("", "Unknown")
    df["unit_price"] = (
        df["unit_price"].astype(str).str.replace("€", "").str.replace(",", ".").str.strip().astype(float)
    )
    df["qty"] = pd.to_numeric(df["qty"], errors="coerce")

    before = len(df)
    df = df[(df["qty"] > 0) & df["date"].notna()]
    stats["invalid_removed"] = before - len(df)

    df = df.assign(total=df["qty"] * df["unit_price"]).sort_values("date")
    stats["rows_out"] = len(df)
    return df, stats


def build_tables(df):
    monthly = (
        df.groupby(df["date"].dt.to_period("M"))
        .agg(orders=("order_id", "count"), units=("qty", "sum"), revenue=("total", "sum"))
        .reset_index()
    )
    monthly["date"] = monthly["date"].dt.strftime("%b %Y")
    monthly = monthly.rename(columns={"date": "month"})

    by_product = (
        df.groupby("product").agg(units=("qty", "sum"), revenue=("total", "sum"))
        .sort_values("revenue", ascending=False).reset_index()
    )
    by_seller = (
        df.groupby("seller").agg(orders=("order_id", "count"), revenue=("total", "sum"))
        .sort_values("revenue", ascending=False).reset_index()
    )
    return monthly, by_product, by_seller


def style_table(ws, start_row, money_cols=()):
    thin = Side(style="thin", color="D0D7DE")
    for cell in ws[start_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="center")
    for i, row in enumerate(ws.iter_rows(min_row=start_row + 1, max_row=ws.max_row), 1):
        for cell in row:
            cell.border = Border(bottom=thin)
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
            if cell.column_letter in money_cols:
                cell.number_format = MONEY
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 4
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width, 12), 40)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def write_report(df, stats, out):
    monthly, by_product, by_seller = build_tables(df)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        pd.DataFrame().to_excel(xw, sheet_name="Summary")
        monthly.to_excel(xw, sheet_name="Monthly", index=False)
        by_product.to_excel(xw, sheet_name="By product", index=False)
        by_seller.to_excel(xw, sheet_name="By seller", index=False)
        df.assign(date=df["date"].dt.date).to_excel(xw, sheet_name="Clean data", index=False)

    wb = load_workbook(out)

    ws = wb["Summary"]
    ws["A1"] = "Sales Report"
    ws["A1"].font = Font(size=22, bold=True, color=DARK)
    ws["A2"] = f"{df['date'].min():%d %b %Y} – {df['date'].max():%d %b %Y}"
    ws["A2"].font = Font(italic=True, color="666666")
    kpis = [
        ("Total revenue", df["total"].sum(), MONEY),
        ("Orders", len(df), "#,##0"),
        ("Units sold", df["qty"].sum(), "#,##0"),
        ("Average order", df["total"].mean(), MONEY),
        ("Best product", by_product.iloc[0]["product"], None),
        ("Best seller", by_seller.iloc[0]["seller"], None),
    ]
    for i, (label, value, fmt) in enumerate(kpis, 4):
        ws.cell(i, 1, label).font = Font(bold=True)
        c = ws.cell(i, 2, value)
        c.font = Font(size=13, bold=True, color=GREEN)
        if fmt:
            c.number_format = fmt
    ws["A11"] = "Data cleaning log"
    ws["A11"].font = Font(size=14, bold=True)
    log = [
        ("Rows received", stats["rows_in"]),
        ("Empty rows removed", stats["empty_removed"]),
        ("Duplicates removed", stats["duplicates_removed"]),
        ("Invalid rows removed", stats["invalid_removed"]),
        ("Clean rows", stats["rows_out"]),
    ]
    for i, (label, value) in enumerate(log, 12):
        ws.cell(i, 1, label)
        ws.cell(i, 2, value)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 24

    ws = wb["Monthly"]
    style_table(ws, 1, money_cols=("D",))
    chart = LineChart()
    chart.title, chart.y_axis.title, chart.height, chart.width = "Revenue per month", "€", 9, 18
    chart.add_data(Reference(ws, min_col=4, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    chart.series[0].graphicalProperties.line.solidFill = GREEN
    chart.series[0].graphicalProperties.line.width = 32000
    chart.series[0].smooth = False
    ws.add_chart(chart, "F2")

    ws = wb["By product"]
    style_table(ws, 1, money_cols=("C",))
    chart = BarChart()
    chart.type, chart.title, chart.height, chart.width = "bar", "Revenue by product", 9, 18
    chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    chart.series[0].graphicalProperties.solidFill = GREEN
    chart.legend = None
    ws.add_chart(chart, "E2")

    style_table(wb["By seller"], 1, money_cols=("C",))
    style_table(wb["Clean data"], 1, money_cols=("G", "H"))
    wb.save(out)


def send_email(to, attachment):
    msg = EmailMessage()
    msg["Subject"] = "Automatic sales report"
    msg["From"] = os.environ["SMTP_USER"]
    msg["To"] = to
    msg.set_content("Hi,\n\nAttached is this week's sales report.\n\n— Your report bot")
    with open(attachment, "rb") as f:
        msg.add_attachment(f.read(), maintype="application", subtype="octet-stream",
                           filename=os.path.basename(attachment))
    with smtplib.SMTP_SSL(os.environ.get("SMTP_HOST", "smtp.gmail.com"), 465) as s:
        s.login(os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"])
        s.send_message(msg)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("csv")
    p.add_argument("-o", "--output", default="sales_report.xlsx")
    p.add_argument("--email", help="send the report to this address")
    args = p.parse_args()

    df, stats = clean(pd.read_csv(args.csv, dtype=str))
    write_report(df, stats, args.output)
    print(f"[ OK ] {stats['rows_in']} rows in → {stats['rows_out']} clean rows")
    print(f"[ OK ] removed: {stats['empty_removed']} empty, {stats['duplicates_removed']} duplicates, "
          f"{stats['invalid_removed']} invalid")
    print(f"[ OK ] report saved to {args.output}")
    if args.email:
        send_email(args.email, args.output)
        print(f"[ OK ] email sent to {args.email}")


if __name__ == "__main__":
    main()
